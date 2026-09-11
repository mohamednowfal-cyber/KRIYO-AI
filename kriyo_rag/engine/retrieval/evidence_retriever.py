"""
Evidence Claims Registry Retriever.
Queries KRIYO_Evidence_Registry.xlsx (500 claims, EV00001 to EV00500).
Maps claims to evidence_id, source_id, document_id, section_id.
"""

import os
from typing import List, Dict, Any, Optional
import openpyxl  # type: ignore
from ...config.settings import settings


class EvidenceRetriever:
    def __init__(self, evidence_path: Optional[str] = None):
        self.evidence_path = str(evidence_path or settings.EVIDENCE_PATH)
        self.claims: List[Dict[str, Any]] = []
        self._load_evidence_registry()

    def _load_evidence_registry(self):
        """Loads Evidence Registry into memory."""
        if not os.path.exists(self.evidence_path):
            return

        try:
            wb = openpyxl.load_workbook(self.evidence_path, data_only=True)
            if "Evidence_Registry" in wb.sheetnames:
                sheet = wb["Evidence_Registry"]
                max_col = sheet.max_column or 0
                max_r = sheet.max_row or 0
                headers = [sheet.cell(1, c).value for c in range(1, max_col + 1)]
                for r in range(2, max_r + 1):
                    row_dict = {}
                    for c, h in enumerate(headers, 1):
                        if h:
                            row_dict[str(h).strip()] = sheet.cell(r, c).value
                    if row_dict.get("evidence_id"):
                        self.claims.append(row_dict)
        except Exception as e:
            print(f"[!] Error loading evidence registry: {e}")

    def search_evidence(self, query: str, craft_id: Optional[str] = None, max_results: int = 5) -> List[Dict[str, Any]]:
        """Searches evidence claims matching craft_id or query terms."""
        q_words = [w.lower() for w in query.split() if len(w) > 2]
        scored_claims = []

        for claim in self.claims:
            score = 0
            if craft_id:
                c_id = str(claim.get("craft_id", "")).upper()
                s_id = str(claim.get("source_id", "")).upper()
                d_id = str(claim.get("document_id", "")).upper()
                target = craft_id.upper()
                if c_id == target or target in s_id or target in d_id:
                    score += 5

            claim_text = str(claim.get("claim_text", "")).lower()
            context = str(claim.get("context", "")).lower()
            combined = f"{claim_text} {context}"

            for w in q_words:
                if w in combined:
                    score += 1

            if score > 0:
                scored_claims.append((score, claim))

        scored_claims.sort(key=lambda x: x[0], reverse=True)
        return [c[1] for c in scored_claims[:max_results]]
