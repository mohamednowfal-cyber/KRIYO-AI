"""
Grounded Answer Generator.
Synthesizes strictly grounded, evidence-backed answers using mode-specific answer composition.
Integrates citation building and synthetic provenance notices.
"""

import os
import re
from typing import Dict, Any, List, Optional
import openpyxl  # type: ignore

from .answerability import AnswerabilityGate
from .citation_builder import CitationBuilder
from .provenance_formatter import ProvenanceFormatter, SYNTHETIC_DEMO_NOTICE
from ...config.settings import settings


class GroundedAnswerGenerator:
    def __init__(self, structured_dir: Optional[str] = None):
        self.structured_dir = str(structured_dir or settings.STRUCTURED_DIR)
        self.gate = AnswerabilityGate(
            min_retrieval_confidence=settings.MIN_RETRIEVAL_CONFIDENCE,
            min_evidence_coverage=settings.MIN_EVIDENCE_COVERAGE
        )
        self.craft_master_records = []
        self._load_craft_master()

    def _load_craft_master(self):
        master_file = os.path.join(self.structured_dir, "KRIYO_Craft_Master.xlsx")
        if os.path.exists(master_file):
            try:
                wb = openpyxl.load_workbook(master_file, data_only=True)
                if "KRIYO_Craft_Master" in wb.sheetnames:
                    sheet = wb["KRIYO_Craft_Master"]
                    max_col = sheet.max_column or 0
                    max_r = sheet.max_row or 0
                    headers = [sheet.cell(1, c).value for c in range(1, max_col + 1)]
                    for r in range(2, max_r + 1):
                        row_dict = {}
                        for c, h in enumerate(headers, 1):
                            if h:
                                row_dict[str(h).strip()] = sheet.cell(r, c).value
                        if row_dict.get("craft_id"):
                            self.craft_master_records.append(row_dict)
            except Exception:
                pass

    def generate_answer(self, query: str, retrieval_output: Dict[str, Any]) -> str:
        """
        Generates grounded, cited response using mode-specific answer composition.
        """
        # 1. Answerability Gate Check
        gate_res = self.gate.evaluate_answerability(retrieval_output)
        if not gate_res["is_answerable"]:
            return gate_res["refusal_msg"]

        analysis = retrieval_output.get("query_analysis", {})
        chunks = retrieval_output.get("retrieved_chunks", [])
        intent = analysis.get("intent", "attribute_lookup")
        q_lower = query.lower()

        # Build citations map
        valid_citations_map, first_citation = CitationBuilder.build_citations(chunks)

        # MODE 1: STATE_INDEX Mode
        if intent == "state_index":
            target_state = analysis.get("target_state")
            if not target_state:
                states_list = [
                    "Kerala", "Karnataka", "Andhra Pradesh", "Telangana", "Odisha", "West Bengal",
                    "Rajasthan", "Gujarat", "Maharashtra", "Madhya Pradesh", "Uttar Pradesh", "Bihar",
                    "Assam", "Manipur", "Meghalaya", "Nagaland", "Tripura", "Sikkim",
                    "Himachal Pradesh", "Uttarakhand", "Jammu & Kashmir", "Tamil Nadu"
                ]
                for st in states_list:
                    if st.lower() in q_lower:
                        target_state = st
                        break
                if not target_state:
                    target_state = "the state"

            crafts_found = []
            if self.craft_master_records:
                for r in self.craft_master_records:
                    if str(r.get("state", "")).strip().lower() == target_state.lower():
                        cname = r.get("canonical_name")
                        if cname and cname not in crafts_found:
                            crafts_found.append(cname)

            if not crafts_found:
                seen = set()
                for chunk in chunks:
                    cname = chunk.get("craft_name")
                    if cname and cname.lower() not in seen and not cname.startswith("DOC"):
                        seen.add(cname.lower())
                        crafts_found.append(cname)

            if crafts_found:
                if len(crafts_found) == 1:
                    crafts_str = crafts_found[0]
                elif len(crafts_found) == 2:
                    crafts_str = f"{crafts_found[0]} and {crafts_found[1]}"
                else:
                    crafts_str = ", ".join(crafts_found[:-1]) + f", and {crafts_found[-1]}"

                citation = first_citation
                ans = f"In {target_state}, the documented traditional crafts in the KRIYO knowledge base include {crafts_str} {citation}."
                return ProvenanceFormatter.format_provenance(ans)

        # MODE 2: MATERIAL / ATTRIBUTE LOOKUP
        if intent == "attribute_lookup" or analysis.get("target_attribute") == "materials":
            materials_found = []
            for chunk in chunks:
                sec_name = chunk.get("section_name", "").lower()
                text = chunk.get("text", "")
                if "material" in sec_name or "raw" in sec_name:
                    citation = f"[{chunk.get('document_id')}: {chunk.get('section_name')}]"
                    ans = f"According to {citation}, the primary materials utilized include {text.splitlines()[1] if len(text.splitlines()) > 1 else text[:200]}."
                    return ProvenanceFormatter.format_provenance(ans)

        # MODE 3: DIRECT FACTUAL SYNTHESIS
        # Extract direct sentences from top chunks
        synthesized_paragraphs = []
        for chunk in chunks[:3]:
            doc_id = chunk.get("document_id", "DOC")
            sec_name = chunk.get("section_name", "Overview")
            text = chunk.get("text", "").strip()

            # Clean header lines
            lines = [l.strip() for l in text.splitlines() if l.strip() and not l.startswith("#")]
            if lines:
                passage = lines[0]
                if len(passage) < 40 and len(lines) > 1:
                    passage = f"{passage} {lines[1]}"
                citation = f"[{doc_id}: {sec_name}]"
                synthesized_paragraphs.append(f"{passage} {citation}")

        if synthesized_paragraphs:
            ans = "\n\n".join(synthesized_paragraphs)
            return ProvenanceFormatter.format_provenance(ans)

        return ProvenanceFormatter.format_provenance(
            f"Based on the KRIYO knowledge base {first_citation}, the documented details correspond to {query}."
        )
