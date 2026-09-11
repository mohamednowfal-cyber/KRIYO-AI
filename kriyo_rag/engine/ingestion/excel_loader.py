"""
Excel Structured Data Loader.
Loads structured datasets from 01_structured workbooks.
"""

import os
from typing import Dict, List, Any
import openpyxl  # type: ignore


class ExcelLoader:
    def __init__(self, structured_dir: str):
        self.structured_dir = structured_dir

    def load_sheet(self, filename: str, sheet_name: str) -> List[Dict[str, Any]]:
        """
        Loads a single worksheet from an Excel workbook into a list of dict records.
        """
        filepath = os.path.join(self.structured_dir, filename)
        if not os.path.exists(filepath):
            return []

        records = []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            if sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                max_col = sheet.max_column or 0
                max_r = sheet.max_row or 0
                headers = [sheet.cell(1, c).value for c in range(1, max_col + 1)]
                for r in range(2, max_r + 1):
                    row_dict = {}
                    for c, h in enumerate(headers, 1):
                        if h:
                            row_dict[str(h).strip()] = sheet.cell(r, c).value
                    if any(v is not None for v in row_dict.values()):
                        records.append(row_dict)
        except Exception as e:
            print(f"[!] Error loading {filename} [{sheet_name}]: {e}")

        return records

    def load_all_key_workbooks(self) -> Dict[str, Dict[str, Any]]:
        """
        Loads standard KRIYO workbooks into a unified lookup dictionary.
        """
        key_files = [
            ("crafts", "KRIYO_Craft_Master.xlsx", "KRIYO_Craft_Master"),
            ("products", "KRIYO_Craft_Products.xlsx", "Products"),
            ("artisans", "KRIYO_Artisan_Profiles.xlsx", "KRIYO_Artisan_Profiles"),
            ("techniques", "KRIYO_Craft_Techniques.xlsx", "Techniques"),
            ("b2b", "KRIYO_B2B_Buyers.xlsx", "Buyers"),
            ("glossary", "KRIYO_Language_Glossary.xlsx", "Terminology"),
            ("provenance", "KRIYO_Provenance.xlsx", "Provenance"),
            ("market", "KRIYO_Market_Intelligence.xlsx", "Market_Intelligence"),
            ("relationships", "KRIYO_Craft_Product_Relationships.xlsx", "Relationships"),
            ("aliases", "KRIYO_Entity_Aliases.xlsx", "Aliases"),
            ("gazetteer", "KRIYO_Location_Gazetteer.xlsx", "Gazetteer")
        ]

        cache = {}
        for key, fname, sname in key_files:
            fpath = os.path.join(self.structured_dir, fname)
            if os.path.exists(fpath):
                records = self.load_sheet(fname, sname)
                cache[key] = {
                    "filename": fname,
                    "filepath": fpath,
                    "sheet_name": sname,
                    "records": records
                }
        return cache
