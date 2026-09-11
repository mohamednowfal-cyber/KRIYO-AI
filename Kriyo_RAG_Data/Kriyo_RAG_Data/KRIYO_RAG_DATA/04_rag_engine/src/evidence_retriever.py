"""
KRIYO RAG Engine — Evidence Claims Registry Retriever
Queries 03_evidence/KRIYO_Evidence_Registry.xlsx (500 claims, EV00001 to EV00500)
Maps claims to evidence_id, source_id, document_id, section_id.
"""

import os
import openpyxl

EVIDENCE_PATH = r"D:\Kriyo_RAG_Data\KRIYO_RAG_DATA\03_evidence\KRIYO_Evidence_Registry.xlsx"


class EvidenceRetriever:
    def __init__(self, evidence_path: str = EVIDENCE_PATH):
        self.evidence_path = evidence_path
        self.claims = []
        self._load_evidence_registry()

    def _load_evidence_registry(self):
        """Loads Evidence Registry into memory."""
        if not os.path.exists(self.evidence_path):
            return

        try:
            wb = openpyxl.load_workbook(self.evidence_path, data_only=True)
            sheet = wb["Evidence_Registry"]
            headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
            for r in range(2, sheet.max_row + 1):
                row_dict = {}
                for c, h in enumerate(headers, 1):
                    if h:
                        row_dict[h] = sheet.cell(r, c).value
                if row_dict.get("evidence_id"):
                    self.claims.append(row_dict)
        except Exception:
            pass

    def search_evidence(self, query: str, craft_id: str = None, max_results: int = 5) -> list[dict]:
        """Searches evidence claims matching craft_id or query terms."""
        q_words = [w.lower() for w in query.split() if len(w) > 2]
        matching_claims = []

        for row in self.claims:
            ev_id = row.get("evidence_id", "")
            claim_txt = str(row.get("claim_text", ""))
            row_cid = str(row.get("craft_id", ""))
            doc_id = str(row.get("document_id", ""))
            source_id = str(row.get("source_id", ""))
            sec_id = str(row.get("section_id", ""))

            score = 0.0
            if craft_id and row_cid.upper() == craft_id.upper():
                score += 2.0

            text_blob = f"{claim_txt} {row.get('claim_type', '')} {row.get('verification_status', '')}".lower()
            matches = sum(1 for w in q_words if w in text_blob)
            score += matches * 0.5

            if score > 0.0:
                matching_claims.append((score, {
                    "chunk_id": f"EvidenceRegistry#{ev_id}",
                    "evidence_id": ev_id,
                    "source_id": source_id if source_id else "KRIYO_EVIDENCE_REGISTRY",
                    "document_id": doc_id if doc_id else "KRIYO_Evidence_Registry.xlsx",
                    "section_id": sec_id if sec_id else "Evidence_Registry",
                    "title": f"Evidence Claim {ev_id}",
                    "section_name": "Evidence Claims",
                    "craft_id": row_cid,
                    "data_status": "synthetic_demo",
                    "folder_type": "evidence",
                    "text": f"Claim ({ev_id}): {claim_txt} [Type: {row.get('claim_type', '')}, Status: {row.get('verification_status', '')}]",
                    "score": round(min(0.95, 0.5 + score * 0.1), 4)
                }))

        matching_claims.sort(key=lambda x: x[0], reverse=True)
        return [item[1] for item in matching_claims[:max_results]]
